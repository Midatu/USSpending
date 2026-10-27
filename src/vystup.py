"""Výstupy pro web (site/) a Excel.

site/meta.json            stav zdrojů, seznam měsíců s počty záznamů, archivy Excelu
site/data/RRRR-MM.json    řádky kontraktů po měsících (web načítá jen zvolené období)
site/tickery.json         directory tickerů
site/kontrakty.xlsx       posledních 12 měsíců (+ archivy kontrakty-RRRR.xlsx za starší roky)
site/zpravy.json          upozornění a přehled sentimentu podle firem (podstránka Zprávy)
site/zpravy/RRRR-MM.json  zprávy po měsících
"""
import datetime as dt
import re
import shutil
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import ceny
from .common import DATA, SITE, WEB, castka_zaznamu, iso_ted, log, uloz_json

BLOKY = "▁▂▃▄▅▆▇█"
WEB_POLE_VYNECHAT = {"id", "cena_datum", "cena_nyni_datum", "prvni_zaznam", "odkaz", "dni"}


VELKA = {"LLC", "LLP", "LP", "JV", "USA", "US", "U.S.", "IT", "IBM", "KBR", "SAIC", "CACI", "BAE", "AT&T", "GE",
         "NASA", "DOD", "DOW", "VA", "FY", "IDIQ", "IDV", "BPA", "GSA", "II", "III", "IV", "PLC", "AG", "SA", "NV",
         "HVAC", "USNS", "USS", "MRO", "O&M", "R&D", "AI", "UAS", "SLA", "CIO", "NAVSEA", "NAVAIR", "SOCOM", "DLA",
         "DISA", "DHS", "FBI", "DEA", "IRS", "CDC", "NIH", "FAA", "NOAA", "VHA", "HHS", "DOE", "EPA"}


def _slovo(w: str, prvni: bool, nazev: bool) -> str:
    holy = w.strip(",.;:()/-")
    if not holy or not any(c.isalpha() for c in holy):
        return w
    if holy.upper() in VELKA or (len(holy) <= 3 and not any(c in "AEIOUY" for c in holy.upper())):
        return w.upper()
    if any(c.isdigit() for c in holy):
        return w
    if nazev or prvni:
        return w[:1].upper() + w[1:].lower() if w[:1].isalpha() else w[:2] + w[2:].lower()
    return w.lower()


def _hezky(text: str, nazev: bool = False) -> str:
    """Text psaný celý VELKÝMI písmeny převede na čitelnou podobu (zkratky ponechá velké)."""
    text = re.sub(r"^\s*UPDATE:\s*", "", text or "", flags=re.I).strip()  # prefix z oprav na SAM.gov
    if not text or not text.isupper():
        return text
    slova = text.split()
    return " ".join(_slovo(w, i == 0, nazev) for i, w in enumerate(slova))


def _popis(z: dict) -> str:
    if z.get("popis_dod"):
        return z["popis_dod"][:1].upper() + z["popis_dod"][1:]
    kandidati = [z.get("popis_award") or "", z.get("popis_usa") or "", z.get("popis_sam") or "", z.get("popis_8k") or ""]
    dlouhe = [k for k in kandidati if len(k) >= 25]
    return _hezky((dlouhe or sorted(kandidati, key=len, reverse=True))[0])


ZADAVATELE = {
    "department of the army": "U.S. Army", "dept of the army": "U.S. Army",
    "department of the navy": "U.S. Navy", "dept of the navy": "U.S. Navy",
    "department of the air force": "U.S. Air Force", "dept of the air force": "U.S. Air Force",
    "department of defense": "Department of Defense", "dept of defense": "Department of Defense",
}


def zadavatel(text: str) -> str:
    text = (text or "").strip()
    if text.isupper():
        text = text.title().replace(" Of ", " of ").replace(" The ", " the ").replace(" And ", " and ")
    # "Veterans Affairs, Department Of" -> "Department of Veterans Affairs"
    m = re.match(r"^(.+?),\s*(Department|Dept\.?|Office|Agency|Bureau)\s+of\s*(the)?$", text, re.I)
    if m:
        text = f"{m.group(2).title()} of {'the ' if m.group(3) else ''}{m.group(1)}"
    text = text.replace("Dept of ", "Department of ").replace("Dept. of ", "Department of ")
    return ZADAVATELE.get(text.lower(), text)


def yahoo(ticker: str) -> str:
    return f"https://finance.yahoo.com/quote/{ticker}"


def text_graf(rada: list) -> str:
    if not rada or len(rada) < 2:
        return ""
    lo, hi = min(rada), max(rada)
    if hi == lo:
        return BLOKY[3] * min(len(rada), 30)
    if len(rada) > 30:
        krok = (len(rada) - 1) / 29
        rada = [rada[round(i * krok)] for i in range(30)]
    return "".join(BLOKY[int((v - lo) / (hi - lo) * 7)] for v in rada)


def priprav(cfg: dict, db: dict, cache_cen: dict) -> list[dict]:
    nove_piid = {z["piid"] for z in db.values() if z["mod"] == "0"}
    radky = []
    for klic, z in db.items():
        r = {
            "id": klic, "datum": z["datum"], "typ": z["typ"],
            "spolecnost": _hezky(z["spolecnost"], nazev=True), "materska": _hezky(z.get("materska", ""), nazev=True),
            "ticker": z.get("ticker", ""), "burza": z.get("burza", ""), "shoda": z.get("shoda", ""),
            "castka": castka_zaznamu(z) or None,
            "popis": _popis(z),
            "poznamka": z.get("poznamka", ""), "agentura": zadavatel(z.get("agentura", "")),
            "zdroje": z.get("zdroje", []),
            "odkaz": (z.get("odkaz_usa") or z.get("odkaz_dod") or z.get("odkaz_sam") or z.get("odkaz_8k") or ""),
            "odkaz_dod": z.get("odkaz_dod", ""), "odkaz_usa": z.get("odkaz_usa", ""),
            "odkaz_sam": z.get("odkaz_sam", ""), "odkaz_8k": z.get("odkaz_8k", ""),
            "piid": z["piid"], "mod": z["mod"], "prvni_zaznam": z.get("prvni_zaznam", ""),
            "puvodni_v_seznamu": z["typ"] == "Navýšení" and z["piid"] in nove_piid,
        }
        if r["ticker"]:
            r.update(ceny.vyhodnot(r["ticker"], r["datum"], cache_cen, int(cfg.get("graf_max_bodu", 40))))
        radky.append(r)
    radky.sort(key=lambda r: (r["datum"], r["castka"] or 0), reverse=True)
    return radky


def _web_radek(r: dict) -> dict:
    """Zkrácený řádek pro web: bez prázdných a nepotřebných polí."""
    w = {k: v for k, v in r.items()
         if k not in WEB_POLE_VYNECHAT and v is not None and v is not False and v != "" and v != []}
    if w.get("graf"):
        w["graf"] = [round(v, 2) for v in w["graf"]]
    return w


# ---------- Excel ----------

def _bunka(ws, hodnota, font=None, fmt=None):
    c = WriteOnlyCell(ws, value=hodnota)
    if font:
        c.font = font
    if fmt:
        c.number_format = fmt
    return c


def _odkaz(url: str, text: str) -> str:
    if not url:
        return ""
    return f'=HYPERLINK("{url.replace(chr(34), "")}","{text.replace(chr(34), "")}")'


TYPY_UPOZORNENI = {
    "negativni_vlna": "Vlna negativních zpráv", "pozitivni_vlna": "Vlna pozitivních zpráv",
    "zhorseni": "Zhoršení sentimentu", "zlepseni": "Zlepšení sentimentu", "zajem": "Nárůst zájmu médií",
}


def typ_upozorneni(typ: str) -> str:
    return "Významný kontrakt" if typ.startswith("kontrakt:") else TYPY_UPOZORNENI.get(typ, typ)


def hodnoceni(s: float) -> str:
    return "pozitivní" if s >= 0.25 else "negativní" if s <= -0.25 else "neutrální"


def excel(radky: list[dict], seznam: dict, stav: dict, cesta, nadpis_obdobi: str, zpravy: dict | None = None):
    wb = Workbook(write_only=True)
    bila, modra = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3A5F")
    odkaz_font = Font(color="0563C1", underline="single")
    zelena, cervena = Font(color="1B7F3B"), Font(color="B3261E")

    def list_s_hlavickou(nazev, sloupce):
        ws = wb.create_sheet(nazev)
        for i, (_, sirka) in enumerate(sloupce):
            ws.column_dimensions[get_column_letter(i + 1)].width = sirka
        ws.freeze_panes = "A2"
        hl = []
        for s, _ in sloupce:
            c = WriteOnlyCell(ws, value=s)
            c.font, c.fill = bila, modra
            c.alignment = Alignment(vertical="center", wrap_text=True)
            hl.append(c)
        ws.append(hl)
        return ws

    sloupce = [
        ("Datum", 12), ("Typ", 12), ("Společnost", 32), ("Ticker", 10), ("Výše ($)", 16),
        ("Popis", 60), ("Zadavatel", 26), ("Zhodnocení", 11), ("Zhodnocení p.a.", 12), ("Cena start", 12),
        ("Cena dnes", 12), ("Měna", 7), ("Vývoj", 32), ("Detail", 12), ("Mateřská společnost", 28),
        ("Burza", 12), ("Shoda tickeru", 14), ("Navazuje na kontrakt v seznamu", 12), ("Zdroje", 22),
        ("Číslo kontraktu", 18),
        ("Modifikace", 11), ("DoD", 8), ("SAM.gov", 9), ("SEC 8-K", 9), ("Poprvé zachyceno (UTC)", 20),
    ]
    ws = list_s_hlavickou("Kontrakty", sloupce)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(sloupce))}{len(radky) + 1}"
    for r in radky:
        zmena = r.get("zmena_pct")
        barva = None if zmena is None else (zelena if zmena >= 0 else cervena)
        ws.append([
            r["datum"], r["typ"], r["spolecnost"],
            _bunka(ws, _odkaz(yahoo(r["ticker"]), r["ticker"]), odkaz_font) if r["ticker"] else "",
            _bunka(ws, r["castka"], fmt="#,##0"),
            r["popis"] + (f" ({r['poznamka']})" if r["poznamka"] else ""), r["agentura"],
            _bunka(ws, (zmena / 100) if zmena is not None else None, barva, "+0.0%;-0.0%;0.0%"),
            _bunka(ws, (r["zmena_pa"] / 100) if r.get("zmena_pa") is not None else None, barva, "+0.0%;-0.0%;0.0%"),
            _bunka(ws, r.get("cena_pri"), fmt="#,##0.00"), _bunka(ws, r.get("cena_nyni"), fmt="#,##0.00"),
            r.get("mena", ""), _bunka(ws, text_graf(r.get("graf") or []), barva),
            _bunka(ws, _odkaz(r["odkaz"], "Otevřít"), odkaz_font),
            r["materska"], r["burza"], r["shoda"],
            ("ano" if r["puvodni_v_seznamu"] else "") if r["typ"] == "Navýšení" else "",
            ", ".join(r["zdroje"]), "" if r["piid"].startswith("8K") else r["piid"],
            "" if r["mod"] == "0" or r["mod"].startswith("?") else r["mod"],
            _bunka(ws, _odkaz(r["odkaz_dod"], "Otevřít"), odkaz_font),
            _bunka(ws, _odkaz(r["odkaz_sam"], "Otevřít"), odkaz_font),
            _bunka(ws, _odkaz(r["odkaz_8k"], "Otevřít"), odkaz_font),
            r["prvni_zaznam"].replace("T", " ").replace("Z", ""),
        ])

    if zpravy is not None:
        clanky = zpravy["clanky_excel"]
        wz = list_s_hlavickou("Zprávy", [("Datum", 12), ("Čas (UTC)", 9), ("Ticker", 12), ("Společnost", 30),
                                         ("Titulek (odkaz na článek)", 90), ("Zdroj", 20), ("Kategorie", 18),
                                         ("Sentiment", 10), ("Hodnocení", 11)])
        wz.auto_filter.ref = f"A1:I{len(clanky) + 1}"
        for z in clanky:
            titulek = WriteOnlyCell(wz, value=z["titulek"])
            titulek.hyperlink, titulek.font = z["url"], odkaz_font
            barva = zelena if z["s"] >= 0.25 else cervena if z["s"] <= -0.25 else None
            wz.append([z["cas"][:10], z["cas"][11:16], ", ".join(z["tickery"]), z["firmy"], titulek, z["zdroj"],
                       z["kat"], _bunka(wz, z["s"], barva, "+0.00;-0.00;0.00"), _bunka(wz, hodnoceni(z["s"]), barva)])
        wu = list_s_hlavickou("Upozornění", [("Od", 12), ("Naposledy", 12), ("Stav", 10), ("Typ", 24),
                                             ("Ticker", 12), ("Společnost", 30), ("Popis", 100)])
        for u in zpravy["upozorneni"]:
            barva = cervena if u["smer"] == "neg" else zelena if u["smer"] == "pos" else None
            wu.append([u["od"], u["do"], "aktivní" if u["aktivni"] else "skončilo",
                       _bunka(wu, typ_upozorneni(u["typ"]), barva), u["ticker"], u["firma"], u["text"]])

    wt = list_s_hlavickou("Tickery", [("Ticker", 12), ("Společnost", 50), ("Burza", 18), ("CIK", 12), ("Zdroj", 14)])
    vse = seznam.get("seznam", [])
    wt.auto_filter.ref = f"A1:E{len(vse) + 1}"
    for t in vse:
        wt.append([_bunka(wt, _odkaz(yahoo(t["ticker"]), t["ticker"]), odkaz_font),
                   t.get("nazev", ""), t.get("burza", ""), t.get("cik"), t.get("zdroj", "")])

    wi = list_s_hlavickou("Stav", [("Položka", 28), ("Hodnota", 90)])
    wi.append(["Období v tomto souboru", nadpis_obdobi])
    wi.append(["Poslední běh (UTC)", stav.get("posledni_beh", "")])
    wi.append(["Seznam tickerů aktualizován", seznam.get("aktualizace", "")])
    for nazev, s in stav.get("zdroje", {}).items():
        if s.get("ok") is None:
            wi.append([nazev, "NEAKTIVNÍ: " + s.get("info", "")])
        else:
            wi.append([nazev, ("OK " + s.get("cas", "")) if s.get("ok") else ("CHYBA: " + s.get("chyba", ""))])
    wb.save(cesta)


# ---------- sestavení webu ----------

def priprav_zpravy(zd: dict) -> dict:
    """Zprávy pro web a Excel: seznam článků, přehled podle firem (7 a 30 dní), upozornění."""
    from .zpravy import NEGATIVNI, POZITIVNI, okno
    firmy = zd["firmy"]
    nazev = {t: _hezky(f.nazev, nazev=True) for t, f in firmy.items()}
    # záznamy bez titulku (jen pro statistiku) se nezobrazují, sentiment a trendy je ale zahrnují
    clanky = sorted((z for z in zd["clanky"].values() if z.get("titulek")), key=lambda z: z["cas"], reverse=True)
    for z in clanky:
        z["firmy"] = ", ".join(nazev.get(t, t) for t in z["tickery"])
    dnes = dt.date.today()
    d = lambda n: (dnes - dt.timedelta(days=n)).isoformat()  # noqa: E731
    prehled = []
    for t in firmy:
        dny = zd["denni"].get(t, {})
        n7, s7, neg7, pos7 = okno(dny, d(6), d(0))
        n30, s30, _, _ = okno(dny, d(36), d(7))
        trend = []
        for i in range(29, -1, -1):
            v = dny.get(d(i))
            trend.append([v[0], round(v[1] / v[0], 2)] if v else None)
        posledni = next((z["cas"] for z in clanky if t in z["tickery"]), "")
        prehled.append({"t": t, "nazev": nazev[t], "n7": n7, "s7": s7, "neg7": neg7, "pos7": pos7,
                        "n30": n30, "s30": s30, "trend": trend, "posledni": posledni})
    prehled.sort(key=lambda r: (r["n7"], r["n30"]), reverse=True)
    upozorneni = sorted(zd["upozorneni"].items(), key=lambda kv: (kv[1]["aktivni"], kv[1]["do"], kv[1]["od"]), reverse=True)
    for uid, u in upozorneni:
        u["id"], u["firma"] = uid, nazev.get(u["ticker"], _hezky(u.get("firma", ""), nazev=True))
    return {"clanky": clanky, "prehled": prehled, "upozorneni": [u for _, u in upozorneni],
            "hranice": [NEGATIVNI, POZITIVNI]}


def sestav(cfg: dict, db: dict, cache_cen: dict, seznam: dict, stav: dict, zpravy_data: dict | None = None):
    """Sestaví web a Excel. Do stav["excel_archivy"] zapíše, kdy byly obnoveny roční archivy (uloží volající)."""
    if SITE.exists():
        shutil.rmtree(SITE)
    (SITE / "data").mkdir(parents=True)
    radky = priprav(cfg, db, cache_cen)
    zpravy = None
    if zpravy_data is not None:
        try:
            zpravy = priprav_zpravy(zpravy_data)
            dni_excel = int(cfg.get("zpravy_excel_dni", 90))
            hranice = (dt.date.today() - dt.timedelta(days=dni_excel)).isoformat()
            zpravy["clanky_excel"] = [z for z in zpravy["clanky"] if z["cas"][:10] >= hranice]
        except Exception as e:  # noqa: BLE001
            log.warning("Výstup zpráv selhal: %s", e)
            zpravy = None

    po_mesicich = defaultdict(list)
    for r in radky:
        po_mesicich[r["datum"][:7]].append(_web_radek(r))
    for mesic, seznam_radku in po_mesicich.items():
        uloz_json(SITE / "data" / f"{mesic}.json", seznam_radku, kompaktni=True)

    # Excel: posledních 12 měsíců + roční archivy za starší roky (obnovují se jednou denně)
    dnes = dt.date.today()
    rok_zpet = (dnes - dt.timedelta(days=365)).isoformat()
    excel([r for r in radky if r["datum"] >= rok_zpet], seznam, stav, SITE / "kontrakty.xlsx",
          f"posledních 12 měsíců (od {rok_zpet})", zpravy)
    archiv_slozka = DATA / "excel"
    archiv_slozka.mkdir(parents=True, exist_ok=True)
    archivy = stav.setdefault("excel_archivy", {})
    roky = sorted({r["datum"][:4] for r in radky if r["datum"][:4] < str(dnes.year)}, reverse=True)
    for rok in roky:
        soubor = archiv_slozka / f"kontrakty-{rok}.xlsx"
        if not soubor.exists() or archivy.get(rok) != dnes.isoformat():
            excel([r for r in radky if r["datum"].startswith(rok)], seznam, stav, soubor, f"rok {rok}")
            archivy[rok] = dnes.isoformat()
            log.info("Excel archiv %s obnoven", rok)
        shutil.copy(soubor, SITE / soubor.name)

    uloz_json(SITE / "meta.json", {
        "aktualizace": iso_ted(), "posledni_beh": stav.get("posledni_beh"),
        "tickery_aktualizace": seznam.get("aktualizace"), "zdroje": stav.get("zdroje", {}),
        "min_castka": cfg.get("min_castka"),
        "mesice": {m: len(v) for m, v in sorted(po_mesicich.items())},
        "excel_archivy": roky,
    }, kompaktni=True)
    if zpravy is not None:
        po_mes = defaultdict(list)
        for z in zpravy["clanky"]:
            po_mes[z["cas"][:7]].append({"cas": z["cas"], "t": z["tickery"], "ti": z["titulek"], "u": z["url"],
                                         "z": z["zdroj"], "s": z["s"], "k": z["kat"], "f": z["firmy"]})
        (SITE / "zpravy").mkdir()
        for mesic, seznam_z in po_mes.items():
            uloz_json(SITE / "zpravy" / f"{mesic}.json", seznam_z, kompaktni=True)
        uloz_json(SITE / "zpravy.json", {
            "aktualizace": iso_ted(), "stav": stav.get("zpravy", {}),
            "mesice": {m: len(v) for m, v in sorted(po_mes.items())},
            "upozorneni": [{k: u.get(k) for k in ("id", "typ", "ticker", "firma", "od", "do", "aktivni", "smer", "text")}
                           | {"nazev_typu": typ_upozorneni(u["typ"])} for u in zpravy["upozorneni"]],
            "prehled": zpravy["prehled"], "hranice": zpravy["hranice"],
            "mesicu_kontraktu": cfg.get("zpravy_mesicu", 6),
        }, kompaktni=True)
    uloz_json(SITE / "tickery.json", [[t["ticker"], t.get("nazev", ""), t.get("burza", "")]
                                      for t in seznam.get("seznam", [])], kompaktni=True)
    for soubor in WEB.iterdir():
        if soubor.is_file():
            shutil.copy(soubor, SITE / soubor.name)
    (SITE / ".nojekyll").write_text("")
    log.info("Výstup sestaven: %d kontraktů v %d měsících", len(radky), len(po_mesicich))
